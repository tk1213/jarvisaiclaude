import io
from datetime import date

import httpx
import pytest
from openpyxl import load_workbook

from app.db import AccountSession
from app.integrations import flowaccount as fa
from app.services import accounting as acc


@pytest.fixture
def db():
    s = AccountSession()
    yield s
    s.close()


def add(db, kind, day, base, **kw):
    data = {"kind": kind, "date": day, "party": kw.pop("party", "บริษัท ทดสอบ จำกัด"), "party_tax_id": "0105555555555", "base": base} | kw
    return acc.save_manual(db, data, created_by="tk")


def test_money_is_exact():
    assert acc.to_satang("1,234.565") == 123457 and acc.to_satang(0.1) == 10
    assert acc.vat_of(acc.to_satang(1234.5)) == 8642  # 86.415 -> 86.42
    with pytest.raises(acc.AccountingError):
        acc.to_satang("abc")


def test_manual_entries_default_to_standard_vat_and_withholding(db):
    e = add(db, "expense", "2026-10-05", "3,000", wht_rate=3, party_tax_id="0105-666-666-666")
    assert (e.vat_satang, e.wht_satang, e.party_tax_id, e.full_tax_invoice) == (21000, 9000, "0105666666666", True)
    no_vat = add(db, "expense", "2026-10-06", 500, vat=0)
    assert no_vat.vat_satang == 0 and no_vat.full_tax_invoice is False  # nothing to claim without VAT
    for bad in ({"base": 0}, {"party": " "}, {"day": "2026-13-01"}, {"party_tax_id": "123"}, {"wht_rate": 20}):
        with pytest.raises(acc.AccountingError):
            add(db, "income", **({"day": "2026-10-01", "base": 100} | bad))
    # Claiming input VAT needs the seller's tax ID (a full tax invoice always has one).
    with pytest.raises(acc.AccountingError, match="เลขผู้เสียภาษี"):
        add(db, "expense", "2026-10-07", 100, party_tax_id="")
    assert add(db, "expense", "2026-10-07", 100, party_tax_id="", full_tax_invoice=False).vat_satang == 700


def test_monthly_summary(db):
    add(db, "income", "2026-10-02", 10000, wht_rate=3)  # customer withheld 300
    add(db, "expense", "2026-10-03", 3000, wht_rate=3)  # we withheld 90 from a company -> ภ.ง.ด.53
    add(db, "expense", "2026-10-04", 1000, wht_rate=5, party_type="person", party="นาย ก")  # 50 -> ภ.ง.ด.3
    add(db, "expense", "2026-10-05", 1000, full_tax_invoice=False)  # abbreviated receipt: its 70 VAT can't be claimed
    voided = add(db, "income", "2026-10-06", 99999)
    voided.excluded = True
    db.commit()
    add(db, "income", "2026-11-01", 5000)  # another month

    s = acc.summary(db, "2026-10")
    assert (s["sales"], s["output_vat"], s["purchases"], s["input_vat"], s["unclaimable_vat"]) == (10000, 700, 5000, 280, 70)
    assert (s["vat_payable"], s["credit_carried_forward"], s["credit_brought_forward"]) == (420, 0, 0)
    assert (s["wht_pnd3"], s["wht_pnd53"], s["wht_credit"]) == (50, 90, 300)
    assert s["profit"] == 10000 - 5000 - 70 and (s["income_count"], s["expense_count"]) == (1, 3)
    assert s["due"] == {"pp30": "2026-11-15", "pp30_efiling": "2026-11-23", "pnd": "2026-11-07", "pnd_efiling": "2026-11-15"}
    assert acc.due_dates("2026-12")["pp30"] == "2027-01-15"


def test_excess_input_vat_is_carried_forward(db):
    add(db, "expense", "2026-08-10", 20000)  # input 1,400, no sales: credit 1,400
    add(db, "income", "2026-09-10", 10000)  # output 700 - credit 1,400: still 700 credit
    add(db, "income", "2026-10-10", 20000)  # output 1,400 - credit 700: pay 700
    sep, oct_ = acc.summary(db, "2026-09"), acc.summary(db, "2026-10")
    assert (sep["credit_brought_forward"], sep["vat_payable"], sep["credit_carried_forward"]) == (1400, 0, 700)
    assert (oct_["credit_brought_forward"], oct_["vat_payable"], oct_["credit_carried_forward"]) == (700, 700, 0)


def test_sync_from_flowaccount_adds_once_and_keeps_local_choices(db):
    client = fa.MockFlowAccountClient()
    assert acc.sync_flowaccount(db, client) == {"added": 3, "updated": 0}
    month = date.today().strftime("%Y-%m")
    income = acc.list_entries(db, "income", month)
    expense = acc.list_entries(db, "expense", month)
    assert {e.doc_no for e in income} == {"IV-MOCK-0001", "CA-MOCK-0001"}
    iv = next(e for e in income if e.doc_no == "IV-MOCK-0001")
    assert (iv.base_satang, iv.vat_satang, iv.wht_satang, iv.source) == (1000000, 70000, 30000, "flowaccount")
    exp = expense[0]
    assert (exp.party_type, exp.wht_satang, exp.full_tax_invoice, exp.description) == ("company", 9000, True, "ค่าบริการขนส่ง")

    acc.update_imported(db, exp, {"party_type": "person", "category": "ค่าขนส่ง", "base_satang": 1})  # amounts can't be changed
    assert acc.sync_flowaccount(db, client) == {"added": 0, "updated": 3}
    db.refresh(exp)
    assert (exp.party_type, exp.category, exp.base_satang) == ("person", "ค่าขนส่ง", 300000)
    assert acc.last_sync(db) is not None
    with pytest.raises(acc.AccountingError):
        acc.delete_entry(db, exp)


def test_voided_documents_are_not_counted(db):
    class Client:
        def list_documents(self, path):
            if path != "/tax-invoices":
                return []
            return [{"recordId": 1, "documentSerial": "IV1", "publishedOn": "2026-10-01", "contactName": "A", "isVat": True, "vatAmount": 70, "grandTotal": 1070, "statusString": "Void"},
                    {"recordId": 2, "documentSerial": "IV2", "publishedOn": "2026-10-01T00:00:00", "contactName": "B", "isVat": False, "grandTotal": 500, "status": 3},
                    {"recordId": 3, "publishedOn": "not a date"}]
    acc.sync_flowaccount(db, Client())
    rows = {e.doc_no: e for e in acc.list_entries(db, "income", "2026-10")}
    assert rows["IV1"].excluded is True and rows["IV2"].excluded is False and rows["IV2"].base_satang == 50000
    assert acc.summary(db, "2026-10")["sales"] == 500


def test_live_client_reads_every_page():
    pages = []

    def handler(request):
        if request.url.path.endswith("/token"):
            return httpx.Response(200, json={"access_token": "t", "expires_in": 3600})
        page = int(request.url.params["currentPage"])
        pages.append((request.url.path, page, request.url.params["pageSize"]))
        rows = [{"recordId": i} for i in range(200)] if page == 1 else [{"recordId": 999}]
        return httpx.Response(200, json={"status": True, "data": {"list": rows, "total": 201}})

    client = fa.FlowAccountClient("https://fa.test/v1", "id", "secret", "flowaccount-api", http=httpx.Client(transport=httpx.MockTransport(handler)))
    assert len(client.list_documents("/expenses")) == 201
    assert pages == [("/v1/expenses", 1, "200"), ("/v1/expenses", 2, "200")]


def test_account_api(client, owner_headers):
    month = "2026-10"
    r = client.post("/account/entries", json={"kind": "income", "date": "2026-10-02", "doc_no": "IV001", "party": "บริษัท เอ จำกัด", "party_tax_id": "0105555555555", "base": 10000}, headers=owner_headers)
    assert r.status_code == 201 and r.json()["vat"] == 700 and r.json()["total"] == 10700
    entry_id = r.json()["id"]
    r = client.post("/account/entries", json={"kind": "expense", "date": "2026-10-03", "party": "ร้าน บี", "base": 1000, "category": "ค่าขนส่ง"}, headers=owner_headers)
    assert r.status_code == 400  # claimable VAT without the seller's tax ID
    client.post("/account/entries", json={"kind": "expense", "date": "2026-10-03", "party": "ร้าน บี", "party_tax_id": "0105666666666", "base": 1000, "category": "ค่าขนส่ง"}, headers=owner_headers)

    rows = client.get("/account/entries", params={"kind": "income", "month": month}, headers=owner_headers).json()
    assert [x["doc_no"] for x in rows] == ["IV001"]
    edited = client.put(f"/account/entries/{entry_id}", json=rows[0] | {"base": 20000, "vat": None}, headers=owner_headers).json()
    assert (edited["base"], edited["vat"], edited["kind"]) == (20000, 1400, "income")
    assert client.get("/account/summary", params={"month": month}, headers=owner_headers).json()["vat_payable"] == 1330

    report = client.get("/account/report", params={"month": month, "kind": "income"}, headers=owner_headers)
    assert report.status_code == 200 and "sales-tax-2026-10.xlsx" in report.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(report.content)).active
    values = [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]
    assert ws["A1"].value == "รายงานภาษีขาย" and "IV001" in values and "บริษัท เอ จำกัด" in values and 1400 in values

    synced = client.post("/account/sync", headers=owner_headers).json()
    assert synced["added"] == 3 and synced["last_sync"]
    info = client.get("/account/info", headers=owner_headers).json()
    assert info["last_sync"] and "ค่าขนส่ง" in info["categories"]
    imported = next(x for x in client.get("/account/entries", params={"kind": "expense", "month": date.today().strftime("%Y-%m")}, headers=owner_headers).json() if x["source"] == "flowaccount")
    assert client.put(f"/account/entries/{imported['id']}", json={"excluded": True}, headers=owner_headers).json()["excluded"] is True
    assert client.delete(f"/account/entries/{imported['id']}", headers=owner_headers).status_code == 400
    assert client.delete(f"/account/entries/{entry_id}", headers=owner_headers).status_code == 204

    assert client.get("/account/entries", params={"kind": "income", "month": "2026-13"}, headers=owner_headers).status_code == 400
    assert client.post("/account/entries", json={"kind": "income", "date": "2026-10-02", "party": "x", "base": -5}, headers=owner_headers).status_code == 422


def test_account_is_admin_only(client, owner_headers):
    client.post("/users", json={"username": "staff", "password": "password123"}, headers=owner_headers)
    token = client.post("/auth/login", json={"username": "staff", "password": "password123"}).json()["access_token"]
    staff = {"Authorization": f"Bearer {token}"}
    assert client.get("/account/summary", params={"month": "2026-10"}, headers=staff).status_code == 403
    assert client.post("/account/sync", headers=staff).status_code == 403
    assert client.get("/account/summary", params={"month": "2026-10"}).status_code == 401
