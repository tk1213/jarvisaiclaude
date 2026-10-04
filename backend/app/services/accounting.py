"""The books behind the Account page: income and expense records, the monthly VAT/withholding-tax summary,
the sales and purchase tax reports, and reading both kinds of document back from FlowAccount.

For a VAT-registered company limited (บริษัท จำกัด จด VAT):
- ภ.พ.30: output VAT (sales) minus claimable input VAT (purchases with a full tax invoice); an excess of input VAT
  is a credit carried to the next month. Due the 15th of the next month (e-Filing the 23rd).
- ภ.ง.ด.3 / ภ.ง.ด.53: tax the company withheld when paying individuals / companies. Due the 7th (e-Filing the 15th).
- Tax that customers withheld from the company's income is a credit against the yearly corporate income tax.
"""

import io
import logging
import re
from datetime import date, datetime, timezone
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.account_models import AccountEntry, AccountSetting
from app.config import get_settings
from app.integrations.flowaccount import EXPENSE_PATH, INCOME_PATHS

log = logging.getLogger(__name__)

VAT_RATE = 7
KINDS = ("income", "expense")
EXPENSE_CATEGORIES = [
    "ค่าสินค้า/วัตถุดิบ",
    "ค่าบริการ/ค่าจ้างทำของ",
    "ค่าเช่า",
    "ค่าขนส่ง",
    "ค่าน้ำมัน/เดินทาง",
    "ค่าน้ำ ค่าไฟ ค่าโทรศัพท์",
    "ค่าโฆษณา",
    "เงินเดือน/ค่าแรง",
    "ค่าธรรมเนียม",
    "อื่นๆ",
]
_MONTH = re.compile(r"^(\d{4})-(0[1-9]|1[0-2])$")
# FlowAccount's list shows the document status as text; these mean it no longer counts.
_VOID = re.compile(r"void|cancel|delete|ยกเลิก|ลบ", re.IGNORECASE)
LAST_SYNC = "flowaccount_last_sync"


class AccountingError(ValueError):
    pass


# --- Money --------------------------------------------------------------------------------------------


def to_satang(value) -> int:
    """Baht (number or text, commas allowed) -> whole satang, rounded half up."""
    if value is None or value == "":
        return 0
    try:
        amount = Decimal(str(value).replace(",", "").strip())
    except InvalidOperation:
        raise AccountingError(f"จำนวนเงินไม่ถูกต้อง: {value}") from None
    return int((amount * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def baht(satang: int) -> float:
    return satang / 100


def vat_of(base_satang: int) -> int:
    return int((Decimal(base_satang) * VAT_RATE / 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def wht_of(base_satang: int, rate: int) -> int:
    return int((Decimal(base_satang) * rate / 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


# --- Months -------------------------------------------------------------------------------------------


def parse_month(month: str) -> tuple[int, int]:
    m = _MONTH.match(month or "")
    if not m:
        raise AccountingError("เดือนต้องอยู่ในรูปแบบ YYYY-MM เช่น 2026-10")
    return int(m[1]), int(m[2])


def month_range(month: str) -> tuple[date, date]:
    """First day of the month and first day of the next."""
    year, mon = parse_month(month)
    start = date(year, mon, 1)
    end = date(year + 1, 1, 1) if mon == 12 else date(year, mon + 1, 1)
    return start, end


def _next_month(year: int, mon: int) -> tuple[int, int]:
    return (year + 1, 1) if mon == 12 else (year, mon + 1)


def due_dates(month: str) -> dict:
    """Filing deadlines for a month's taxes (the dates in the law; a weekend or holiday moves them to the next working day)."""
    year, mon = _next_month(*parse_month(month))
    return {
        "pp30": date(year, mon, 15).isoformat(),
        "pp30_efiling": date(year, mon, 23).isoformat(),
        "pnd": date(year, mon, 7).isoformat(),
        "pnd_efiling": date(year, mon, 15).isoformat(),
    }


# --- Entries ------------------------------------------------------------------------------------------


def entry_out(e: AccountEntry) -> dict:
    return {
        "id": e.id,
        "kind": e.kind,
        "date": e.entry_date.isoformat(),
        "doc_no": e.doc_no,
        "party": e.party,
        "party_tax_id": e.party_tax_id,
        "party_branch": e.party_branch,
        "party_type": e.party_type,
        "description": e.description,
        "category": e.category,
        "base": baht(e.base_satang),
        "vat": baht(e.vat_satang),
        "total": baht(e.base_satang + e.vat_satang),
        "wht_rate": e.wht_rate,
        "wht": baht(e.wht_satang),
        "full_tax_invoice": e.full_tax_invoice,
        "excluded": e.excluded,
        "source": e.source,
        "external_status": e.external_status,
        "created_by": e.created_by,
    }


def list_entries(db: Session, kind: str, month: str) -> list[AccountEntry]:
    if kind not in KINDS:
        raise AccountingError("kind ต้องเป็น income หรือ expense")
    start, end = month_range(month)
    return list(
        db.scalars(
            select(AccountEntry)
            .where(AccountEntry.kind == kind, AccountEntry.entry_date >= start, AccountEntry.entry_date < end)
            .order_by(AccountEntry.entry_date, AccountEntry.id)
        )
    )


def save_manual(db: Session, data: dict, created_by: str = "", entry: AccountEntry | None = None) -> AccountEntry:
    """Add (or edit) a record typed on the dashboard. VAT and withholding default to the standard amounts."""
    kind = data.get("kind") or (entry.kind if entry else "")
    if kind not in KINDS:
        raise AccountingError("ต้องเลือกว่าเป็นรายรับหรือรายจ่าย")
    try:
        entry_date = date.fromisoformat(str(data.get("date") or ""))
    except ValueError:
        raise AccountingError("วันที่ไม่ถูกต้อง") from None
    party = (data.get("party") or "").strip()
    if not party:
        raise AccountingError("ต้องใส่ชื่อ" + ("ลูกค้า" if kind == "income" else "ผู้ขาย/ผู้รับเงิน"))
    base = to_satang(data.get("base"))
    if base <= 0:
        raise AccountingError("ยอดก่อน VAT ต้องมากกว่า 0")
    vat = vat_of(base) if data.get("vat") is None else to_satang(data.get("vat"))
    rate = int(data.get("wht_rate") or 0)
    if not 0 <= rate <= 15:
        raise AccountingError("อัตราหัก ณ ที่จ่ายต้องอยู่ระหว่าง 0-15%")
    wht = wht_of(base, rate) if data.get("wht") is None else to_satang(data.get("wht"))
    if vat < 0 or wht < 0:
        raise AccountingError("VAT และภาษีหัก ณ ที่จ่ายต้องไม่ติดลบ")
    tax_id = re.sub(r"\D", "", data.get("party_tax_id") or "")
    if tax_id and len(tax_id) != 13:
        raise AccountingError("เลขผู้เสียภาษีต้องมี 13 หลัก")
    claimable = bool(data.get("full_tax_invoice", True)) and vat > 0
    if kind == "expense" and claimable and not tax_id:
        raise AccountingError("ใบกำกับภาษีเต็มรูปต้องมีเลขผู้เสียภาษีของผู้ขาย ใส่เลข 13 หลัก หรือเอาติ๊ก \"ได้ใบกำกับภาษีเต็มรูป\" ออก (VAT จะไม่นับเป็นภาษีซื้อ)")

    e = entry or AccountEntry(kind=kind, source="manual", created_by=created_by)
    e.entry_date = entry_date
    e.doc_no = (data.get("doc_no") or "").strip()
    e.party = party
    e.party_tax_id = tax_id
    e.party_branch = (data.get("party_branch") or "").strip()
    e.party_type = "person" if data.get("party_type") == "person" else "company"
    e.description = (data.get("description") or "").strip()
    e.category = (data.get("category") or "").strip()
    e.base_satang, e.vat_satang, e.wht_rate, e.wht_satang = base, vat, rate, wht
    e.full_tax_invoice = claimable
    e.excluded = bool(data.get("excluded", False))
    db.add(e)
    db.commit()
    return e


def update_imported(db: Session, entry: AccountEntry, data: dict) -> AccountEntry:
    """A record read from FlowAccount: its amounts come from FlowAccount, only the bookkeeping choices are local."""
    if "excluded" in data:
        entry.excluded = bool(data["excluded"])
    if "full_tax_invoice" in data:
        entry.full_tax_invoice = bool(data["full_tax_invoice"]) and entry.vat_satang > 0
    if "category" in data:
        entry.category = (data["category"] or "").strip()
    if data.get("party_type") in ("person", "company"):
        entry.party_type = data["party_type"]
    db.commit()
    return entry


def delete_entry(db: Session, entry: AccountEntry) -> None:
    if entry.source != "manual":
        raise AccountingError("รายการจาก FlowAccount ลบไม่ได้ ให้ติ๊ก \"ไม่นับ\" แทน (หรือลบเอกสารใน FlowAccount)")
    db.delete(entry)
    db.commit()


# --- FlowAccount --------------------------------------------------------------------------------------


def _fa_entry(kind: str, path: str, doc: dict) -> dict | None:
    """A FlowAccount list item as the fields of an AccountEntry (None when it can't be used)."""
    record = doc.get("recordId") or doc.get("documentId") or doc.get("id") or doc.get("documentSerial")
    try:
        entry_date = date.fromisoformat(str(doc.get("publishedOn") or "")[:10])
    except ValueError:
        return None
    if not record:
        return None
    vat = to_satang(doc.get("vatAmount")) if doc.get("isVat", True) else 0
    total = doc.get("grandTotal")
    base = to_satang(total) - vat if total is not None else to_satang(doc.get("totalAfterDiscount") or doc.get("subTotal"))
    rate = int(doc.get("documentWithholdingTaxPercentage") or 0)
    wht = to_satang(doc.get("documentWithholdingTaxAmount")) if doc.get("documentWithholdingTaxAmount") else wht_of(base, rate)
    items = doc.get("items") or []
    first_item = (items[0].get("name") or items[0].get("description") or "") if items and isinstance(items[0], dict) else ""
    status = str(doc.get("statusString") or doc.get("status") or "")
    tax_id = re.sub(r"\D", "", str(doc.get("contactTaxId") or ""))
    return {
        "kind": kind,
        "external_id": f"{path.strip('/')}:{record}",
        "entry_date": entry_date,
        "doc_no": str(doc.get("documentSerial") or ""),
        "party": str(doc.get("contactName") or ""),
        "party_tax_id": tax_id,
        "party_branch": str(doc.get("contactBranch") or ""),
        "party_type": "person" if int(doc.get("contactGroup") or 3) == 1 else "company",
        "description": str(doc.get("remarks") or first_item or "").strip()[:500],
        "base_satang": base,
        "vat_satang": vat,
        "wht_rate": rate,
        "wht_satang": wht,
        "external_status": status[:64],
        "void": bool(_VOID.search(status)),
        "full_tax_invoice": vat > 0 and len(tax_id) == 13,
    }


def sync_flowaccount(db: Session, client) -> dict:
    """Read sales tax invoices and expenses from FlowAccount; new ones are added, known ones refreshed."""
    counts = {"added": 0, "updated": 0}
    jobs = [("income", path) for path in INCOME_PATHS] + [("expense", EXPENSE_PATH)]
    existing = {
        (e.kind, e.external_id): e for e in db.scalars(select(AccountEntry).where(AccountEntry.source == "flowaccount"))
    }
    for kind, path in jobs:
        for doc in client.list_documents(path):
            fields = _fa_entry(kind, path, doc)
            if fields is None:
                continue
            void = fields.pop("void")
            claimable = fields.pop("full_tax_invoice")
            entry = existing.get((kind, fields["external_id"]))
            if entry is None:
                entry = AccountEntry(source="flowaccount", created_by="FlowAccount", full_tax_invoice=claimable, excluded=void)
                existing[(kind, fields["external_id"])] = entry
                db.add(entry)
                counts["added"] += 1
            else:
                fields.pop("party_type")  # may have been corrected on the dashboard
                if entry.description:
                    fields.pop("description")
                if void:
                    entry.excluded = True
                counts["updated"] += 1
            for key, value in fields.items():
                setattr(entry, key, value)
            if entry.vat_satang == 0:
                entry.full_tax_invoice = False
    set_setting(db, LAST_SYNC, datetime.now(timezone.utc).isoformat(timespec="seconds"))
    db.commit()
    return counts


def get_setting(db: Session, key: str) -> str | None:
    row = db.get(AccountSetting, key)
    return row.value if row else None


def set_setting(db: Session, key: str, value: str) -> None:
    row = db.get(AccountSetting, key) or AccountSetting(key=key)
    row.value = value
    db.add(row)


def last_sync(db: Session) -> datetime | None:
    value = get_setting(db, LAST_SYNC)
    return datetime.fromisoformat(value) if value else None


# --- Monthly summary ----------------------------------------------------------------------------------


def _month_totals(db: Session, start: date, end: date) -> dict:
    rows = db.scalars(
        select(AccountEntry).where(AccountEntry.entry_date >= start, AccountEntry.entry_date < end, AccountEntry.excluded.is_(False))
    )
    t = dict.fromkeys(
        ("sales", "output_vat", "purchases", "input_vat", "unclaimable_vat", "wht_pnd3", "wht_pnd53", "wht_credit", "income_count", "expense_count"), 0
    )
    for e in rows:
        if e.kind == "income":
            t["sales"] += e.base_satang
            t["output_vat"] += e.vat_satang
            t["wht_credit"] += e.wht_satang
            t["income_count"] += 1
        else:
            t["purchases"] += e.base_satang
            if e.full_tax_invoice:
                t["input_vat"] += e.vat_satang
            else:
                t["unclaimable_vat"] += e.vat_satang
            t["wht_pnd3" if e.party_type == "person" else "wht_pnd53"] += e.wht_satang
            t["expense_count"] += 1
    return t


def summary(db: Session, month: str) -> dict:
    """What the company owes the Revenue Department for one month, with the VAT credit carried in from earlier months."""
    target = parse_month(month)
    first = db.scalar(select(func.min(AccountEntry.entry_date)).where(AccountEntry.excluded.is_(False)))
    credit = 0  # excess input VAT carried forward (ภาษีซื้อเกินภาษีขายยกมา)
    if first is not None:
        ym = (first.year, first.month)
        while ym < target:
            start, end = month_range(f"{ym[0]:04d}-{ym[1]:02d}")
            t = _month_totals(db, start, end)
            net = t["output_vat"] - t["input_vat"] - credit
            credit = -net if net < 0 else 0
            ym = _next_month(*ym)
    start, end = month_range(month)
    t = _month_totals(db, start, end)
    net = t["output_vat"] - t["input_vat"] - credit
    # Expenses without a full tax invoice still cost money; their VAT just can't be claimed.
    profit = t["sales"] - t["purchases"] - t["unclaimable_vat"]
    return {
        "month": month,
        "sales": baht(t["sales"]),
        "output_vat": baht(t["output_vat"]),
        "purchases": baht(t["purchases"]),
        "input_vat": baht(t["input_vat"]),
        "unclaimable_vat": baht(t["unclaimable_vat"]),
        "credit_brought_forward": baht(credit),
        "vat_payable": baht(max(net, 0)),
        "credit_carried_forward": baht(max(-net, 0)),
        "wht_pnd3": baht(t["wht_pnd3"]),
        "wht_pnd53": baht(t["wht_pnd53"]),
        "wht_credit": baht(t["wht_credit"]),
        "profit": baht(profit),
        "income_count": t["income_count"],
        "expense_count": t["expense_count"],
        "due": due_dates(month),
    }


# --- Tax reports (Excel) ------------------------------------------------------------------------------

_THAI_MONTHS = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]


def tax_report_xlsx(db: Session, month: str, kind: str) -> bytes:
    """รายงานภาษีขาย (income) or รายงานภาษีซื้อ (expense) for one month, in the columns the Revenue Department expects."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    year, mon = parse_month(month)
    s = get_settings()
    title = "รายงานภาษีขาย" if kind == "income" else "รายงานภาษีซื้อ"
    entries = [e for e in list_entries(db, kind, month) if not e.excluded and (kind == "income" or e.full_tax_invoice)]

    wb = Workbook()
    ws = wb.active
    ws.title = title
    bold = Font(bold=True)
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"เดือนภาษี {_THAI_MONTHS[mon - 1]} พ.ศ. {year + 543}"])
    ws.append([f"ชื่อผู้ประกอบการ {s.company_name or '-'}", "", "", f"เลขประจำตัวผู้เสียภาษีอากร {s.company_tax_id or '-'}", "", f"สถานประกอบการ {s.company_branch or '-'}"])
    ws.append([])
    party = "ชื่อผู้ซื้อสินค้า/ผู้รับบริการ" if kind == "income" else "ชื่อผู้ขายสินค้า/ผู้ให้บริการ"
    header = ["ลำดับ", "วัน เดือน ปี", "เลขที่ใบกำกับภาษี", party, "เลขประจำตัวผู้เสียภาษีอากร", "สถานประกอบการ", "มูลค่าสินค้าหรือบริการ", "จำนวนเงินภาษีมูลค่าเพิ่ม"]
    ws.append(header)
    for cell in ws[ws.max_row]:
        cell.font = bold
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    base_total = vat_total = 0
    for n, e in enumerate(entries, 1):
        d = e.entry_date
        ws.append([n, f"{d.day:02d}/{d.month:02d}/{d.year + 543}", e.doc_no, e.party, e.party_tax_id, e.party_branch, baht(e.base_satang), baht(e.vat_satang)])
        base_total += e.base_satang
        vat_total += e.vat_satang
    ws.append(["", "", "", "รวม", "", "", baht(base_total), baht(vat_total)])
    for cell in ws[ws.max_row]:
        cell.font = bold
    for row in ws.iter_rows(min_row=6, min_col=7, max_col=8):
        for cell in row:
            cell.number_format = "#,##0.00"
    for col, width in zip("ABCDEFGH", (7, 13, 20, 40, 22, 16, 20, 20)):
        ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
